"""Generate the Review 1 DOCX, PDF, and PPTX from the maintained Markdown report."""

from __future__ import annotations

import html
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PptxInches, Pt as PptxPt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph,
    Preformatted,
    Image as ReportImage,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.platypus.tableofcontents import TableOfContents

from review_visuals import (
    ARCHITECTURE_OUTPUT,
    RENTAL_SEQUENCE_OUTPUT,
    SYSTEM_CONTEXT_OUTPUT,
    TRUST_BOUNDARY_OUTPUT,
    build_all as build_review_visuals,
)

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "ShaderMesh-Review1-Report.md"
DOCX_OUTPUT = HERE / "ShaderMesh-Review1-Report.docx"
PDF_OUTPUT = HERE / "ShaderMesh-Review1-Report.pdf"
PPTX_OUTPUT = HERE / "ShaderMesh-Review1-Presentation.pptx"


@dataclass
class Block:
    kind: str
    value: object


def parse_report() -> list[Block]:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    start = lines.index("## Abstract")
    lines = lines[start:]
    blocks: list[Block] = []
    i = 0
    skip_toc = False

    while i < len(lines):
        line = lines[i].rstrip()
        if line == "## Table of Contents":
            skip_toc = True
            i += 1
            continue
        if skip_toc:
            if line == "## Abbreviations":
                skip_toc = False
            else:
                i += 1
                continue
        if not line or line == "---":
            i += 1
            continue
        if line.startswith("```"):
            language = line[3:].strip()
            i += 1
            code: list[str] = []
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            blocks.append(Block("code", (language, "\n".join(code))))
            i += 1
            continue
        image = re.fullmatch(r"!\[([^]]+)]\(([^)]+)\)", line)
        if image:
            blocks.append(Block("image", (image.group(1), image.group(2))))
            i += 1
            continue
        if line.startswith("### "):
            blocks.append(Block("heading2", line[4:].strip()))
            i += 1
            continue
        if line.startswith("## "):
            blocks.append(Block("heading1", line[3:].strip()))
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|"):
            raw_rows: list[str] = []
            while i < len(lines) and lines[i].startswith("|"):
                raw_rows.append(lines[i])
                i += 1
            rows = [_split_table_row(row) for row in raw_rows]
            rows = [row for row in rows if not all(re.fullmatch(r":?-{3,}:?", c) for c in row)]
            blocks.append(Block("table", rows))
            continue
        match = re.match(r"^([-*]|\d+\.)\s+(.*)$", line)
        if match:
            marker, text = match.groups()
            i += 1
            while i < len(lines) and lines[i].startswith(("  ", "\t")):
                text += " " + lines[i].strip()
                i += 1
            blocks.append(Block("list", (marker, text)))
            continue
        if line.startswith("> "):
            quote = line[2:].strip()
            i += 1
            while i < len(lines) and lines[i].startswith("> "):
                quote += " " + lines[i][2:].strip()
                i += 1
            blocks.append(Block("quote", quote))
            continue

        paragraph = line.strip()
        i += 1
        while i < len(lines):
            following = lines[i].rstrip()
            if (
                not following
                or following.startswith(("## ", "### ", "```", "|", "> "))
                or re.match(r"^([-*]|\d+\.)\s+", following)
            ):
                break
            paragraph += " " + following.strip()
            i += 1
        blocks.append(Block("paragraph", paragraph))
    return blocks


def _split_table_row(row: str) -> list[str]:
    return [cell.strip() for cell in row.strip().strip("|").split("|")]


def _plain(text: str) -> str:
    return re.sub(r"(\*\*|`|\*)", "", text)


def _docx_inline(paragraph, text: str) -> None:
    token = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")
    cursor = 0
    for match in token.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        raw = match.group(0)
        if raw.startswith("**"):
            run = paragraph.add_run(raw[2:-2])
            run.bold = True
        elif raw.startswith("`"):
            run = paragraph.add_run(raw[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(10)
        else:
            run = paragraph.add_run(raw[1:-1])
            run.italic = True
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def _prevent_table_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def _set_table_cell_margins(cell, top: int = 70, start: int = 80, bottom: int = 70, end: int = 80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_widths(table, widths: list[float]) -> None:
    table.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)


def _set_paragraph_bottom_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    p_pr.append(borders)


def _add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    begin_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    begin_run._r.append(begin)

    instruction_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    instruction_run._r.append(instr)

    separator_run = paragraph.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separator_run._r.append(separate)

    if placeholder:
        paragraph.add_run(placeholder)

    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def _set_style_font(style, name: str) -> None:
    style.font.name = name
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attribute}"), name)


def _configure_docx(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    _set_style_font(normal, "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = Pt(15)
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, before, after in (
        ("Title", 24, 0, 18),
        ("Subtitle", 15, 0, 12),
        ("Heading 1", 14, 18, 8),
        ("Heading 2", 11.5, 12, 5),
    ):
        style = styles[style_name]
        _set_style_font(style, "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    _set_style_font(code, "Courier New")
    code.font.size = Pt(8.5)
    code.paragraph_format.left_indent = Inches(0.25)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.keep_together = True

    settings = document.settings._element
    for existing in settings.findall(qn("w:updateFields")):
        settings.remove(existing)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "false")
    settings.append(update)


def _configure_section(section) -> None:
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.48)
    section.footer_distance = Inches(0.48)


def build_docx(blocks: list[Block]) -> None:
    document = Document()
    _configure_docx(document)
    cover_section = document.sections[0]
    _configure_section(cover_section)

    border = OxmlElement("w:pgBorders")
    border.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "16")
        edge.set(qn("w:space"), "24")
        edge.set(qn("w:color"), "000000")
        border.append(edge)
    cover_section._sectPr.append(border)

    document.add_paragraph("[INSTITUTION NAME]", style="Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = document.add_paragraph("[DEPARTMENT NAME]\n[DEGREE / PROGRAM NAME]")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)

    p = document.add_paragraph("PROJECT REVIEW 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(14)

    p = document.add_paragraph("SHADERMESH", style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(28)
    p = document.add_paragraph("A Decentralized Marketplace for Renting Compute", style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)

    p = document.add_paragraph("A Proposal and Design Report")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    p.paragraph_format.space_after = Pt(24)

    cover_table = document.add_table(rows=5, cols=2)
    cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover_table.style = "Table Grid"
    _set_table_widths(cover_table, [1.35, 4.25])
    cover_values = [
        ("Course", "[COURSE CODE AND COURSE TITLE]"),
        ("Faculty Guide", "[FACULTY GUIDE NAME AND DESIGNATION]"),
        ("Submitted by", "[STUDENT NAME(S) AND ROLL NUMBER(S)]"),
        ("Academic Year", "[2026-2027]"),
        ("Submission Date", "[DD MONTH YYYY]"),
    ]
    for row, values in zip(cover_table.rows, cover_values):
        row.height = Inches(0.42)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        _prevent_table_row_split(row)
        row.cells[0].text, row.cells[1].text = values
        row.cells[0].paragraphs[0].runs[0].bold = True
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_table_cell_margins(cell, top=90, start=110, bottom=90, end=110)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_after = Pt(2)

    body_section = document.add_section(WD_SECTION.NEW_PAGE)
    _configure_section(body_section)
    body_section.header.is_linked_to_previous = False
    body_section.footer.is_linked_to_previous = False
    header = body_section.header.paragraphs[0]
    header.text = "SHADERMESH  |  PROJECT REVIEW 1"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.runs[0].font.name = "Times New Roman"
    header.runs[0].font.size = Pt(8.5)
    _set_paragraph_bottom_border(header)
    footer = body_section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Page ")
    _add_field(footer, "PAGE", "1")
    for run in footer.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(8.5)

    page_numbering = OxmlElement("w:pgNumType")
    page_numbering.set(qn("w:start"), "1")
    body_section._sectPr.append(page_numbering)

    document.add_heading("Table of Contents", level=1)
    toc = document.add_paragraph()
    toc.paragraph_format.line_spacing = 1.15
    _add_field(toc, 'TOC \\o "1-1" \\h \\z \\u', "Right-click and select Update Field.")
    document.add_page_break()

    for block in blocks:
        if block.kind == "heading1":
            title = str(block.value)
            if title == "Table of Contents":
                continue
            if title == "1. Introduction":
                document.add_page_break()
            document.add_heading(_plain(title), level=1)
        elif block.kind == "heading2":
            document.add_heading(_plain(str(block.value)), level=2)
        elif block.kind == "paragraph":
            paragraph = document.add_paragraph()
            _docx_inline(paragraph, str(block.value))
        elif block.kind == "quote":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.35)
            paragraph.paragraph_format.right_indent = Inches(0.35)
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(6)
            _docx_inline(paragraph, str(block.value))
            for run in paragraph.runs:
                run.italic = True
        elif block.kind == "list":
            marker, text = block.value
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.left_indent = Inches(0.3)
            paragraph.paragraph_format.first_line_indent = Inches(-0.2)
            paragraph.paragraph_format.space_after = Pt(3)
            _docx_inline(paragraph, f"{marker} {text}")
        elif block.kind == "code":
            _language, code = block.value
            document.add_paragraph(str(code), style="Code Block")
        elif block.kind == "image":
            caption, path = block.value
            image_path = HERE / str(path)
            picture = document.add_paragraph()
            picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # Exact line spacing clips inline pictures in Word. Single spacing
            # expands the line box to the full image height.
            picture.paragraph_format.line_spacing = 1.0
            picture.paragraph_format.keep_with_next = True
            picture.paragraph_format.keep_together = True
            picture.paragraph_format.space_before = Pt(4)
            picture.paragraph_format.space_after = Pt(2)
            picture.add_run().add_picture(str(image_path), width=Inches(6.55))
            caption_paragraph = document.add_paragraph()
            caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption_paragraph.paragraph_format.line_spacing = 1.0
            caption_paragraph.paragraph_format.keep_together = True
            caption_paragraph.paragraph_format.space_after = Pt(8)
            caption_run = caption_paragraph.add_run(str(caption))
            caption_run.italic = True
            caption_run.font.size = Pt(9)
        elif block.kind == "table":
            rows = block.value
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            headers = rows[0]
            if len(headers) == 2:
                widths = [1.35, 5.35] if headers[0] == "Term" else [2.0, 4.7]
            elif len(headers) == 3:
                widths = [1.25, 2.65, 2.8]
            elif len(headers) == 4:
                widths = [0.8, 2.05, 2.25, 1.6]
            else:
                widths = [6.7 / len(headers)] * len(headers)
            _set_table_widths(table, widths)
            _set_repeat_table_header(table.rows[0])
            for row_index, row_values in enumerate(rows):
                row = table.rows[row_index]
                _prevent_table_row_split(row)
                for column_index, value in enumerate(row_values):
                    cell = table.cell(row_index, column_index)
                    cell.text = ""
                    _set_table_cell_margins(cell)
                    paragraph = cell.paragraphs[0]
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.paragraph_format.line_spacing = 1.0
                    _docx_inline(paragraph, value)
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(8.5)
                        if row_index == 0:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                    if row_index == 0:
                        _set_cell_shading(cell, "000000")
            spacer = document.add_paragraph()
            spacer.paragraph_format.space_after = Pt(0)
            spacer.paragraph_format.line_spacing = Pt(4)

    document.core_properties.title = "ShaderMesh Project Review 1: Proposal and Design"
    document.core_properties.subject = "Decentralized compute rental marketplace"
    document.core_properties.author = "[STUDENT NAME(S)]"
    document.core_properties.keywords = "Ethereum, compute rental, GPU, escrow, staking"
    document.save(DOCX_OUTPUT)


def _pdf_inline(text: str) -> str:
    escaped = html.escape(text)
    escaped = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", escaped)
    escaped = re.sub(r"`([^`]+)`", r'<font name="Courier">\1</font>', escaped)
    escaped = re.sub(r"\*([^*]+)\*", r"<i>\1</i>", escaped)
    return escaped


class ReportTemplate(BaseDocTemplate):
    def __init__(self, filename: Path, **kwargs):
        super().__init__(str(filename), **kwargs)
        frame = Frame(self.leftMargin, self.bottomMargin, self.width, self.height, id="body")
        self.addPageTemplates(PageTemplate(id="report", frames=frame, onPage=self._decorate_page))
        self._bookmark = 0

    def beforeDocument(self) -> None:
        self._bookmark = 0
        super().beforeDocument()

    def _decorate_page(self, canvas, doc) -> None:
        canvas.saveState()
        canvas.setStrokeColor(colors.black)
        if doc.page == 1:
            canvas.setLineWidth(1.2)
            canvas.rect(0.48 * inch, 0.48 * inch, A4[0] - 0.96 * inch, A4[1] - 0.96 * inch)
        else:
            canvas.setFont("Times-Roman", 8)
            canvas.drawString(self.leftMargin, A4[1] - 0.42 * inch, "SHADERMESH  |  PROJECT REVIEW 1")
            canvas.drawCentredString(A4[0] / 2, 0.4 * inch, f"Page {doc.page - 1}")
            canvas.line(self.leftMargin, A4[1] - 0.48 * inch, A4[0] - self.rightMargin, A4[1] - 0.48 * inch)
        canvas.restoreState()

    def afterFlowable(self, flowable) -> None:
        if isinstance(flowable, Paragraph) and flowable.style.name in ("Heading1", "Heading2"):
            text = flowable.getPlainText()
            if text == "Table of Contents" or flowable.style.name != "Heading1":
                return
            level = 0
            key = f"heading-{self._bookmark}"
            self._bookmark += 1
            self.canv.bookmarkPage(key)
            self.canv.addOutlineEntry(text, key, level=level, closed=False)
            self.notify("TOCEntry", (level, text, self.page - 1, key))


def _pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            "ReportBody",
            parent=styles["BodyText"],
            fontName="Times-Roman",
            fontSize=10.5,
            leading=15,
            alignment=TA_JUSTIFY,
            textColor=colors.black,
            spaceAfter=6,
        )
    )
    heading1 = styles["Heading1"]
    heading1.fontName = "Times-Bold"
    heading1.fontSize = 14
    heading1.leading = 17
    heading1.textColor = colors.black
    heading1.spaceBefore = 13
    heading1.spaceAfter = 7
    heading1.keepWithNext = True
    heading2 = styles["Heading2"]
    heading2.fontName = "Times-Bold"
    heading2.fontSize = 11.5
    heading2.leading = 14
    heading2.textColor = colors.black
    heading2.spaceBefore = 9
    heading2.spaceAfter = 4
    heading2.keepWithNext = True
    styles.add(
        ParagraphStyle(
            "List",
            parent=styles["ReportBody"],
            leftIndent=18,
            firstLineIndent=-14,
            spaceAfter=3,
        )
    )
    styles.add(
        ParagraphStyle(
            "Quote",
            parent=styles["ReportBody"],
            leftIndent=24,
            rightIndent=24,
            fontName="Times-Italic",
        )
    )
    return styles


def build_pdf(blocks: list[Block]) -> None:
    styles = _pdf_styles()
    story = [
        Spacer(1, 0.35 * inch),
        Paragraph("[INSTITUTION NAME]", ParagraphStyle("CoverInstitution", fontName="Times-Bold", fontSize=18, leading=22, alignment=TA_CENTER)),
        Paragraph("[DEPARTMENT NAME]", ParagraphStyle("CoverDepartment", fontName="Times-Roman", fontSize=12, leading=18, alignment=TA_CENTER)),
        Paragraph("[DEGREE / PROGRAM NAME]", ParagraphStyle("CoverProgram", fontName="Times-Roman", fontSize=12, leading=18, alignment=TA_CENTER)),
        Spacer(1, 0.75 * inch),
        Paragraph("PROJECT REVIEW 1", ParagraphStyle("CoverReview", fontName="Times-Bold", fontSize=14, leading=18, alignment=TA_CENTER)),
        Spacer(1, 0.28 * inch),
        Paragraph("SHADERMESH", ParagraphStyle("CoverTitle", fontName="Times-Bold", fontSize=28, leading=32, alignment=TA_CENTER)),
        Paragraph("A Decentralized Marketplace for Renting Compute", ParagraphStyle("CoverSubtitle", fontName="Times-Bold", fontSize=15, leading=20, alignment=TA_CENTER)),
        Spacer(1, 0.35 * inch),
        Paragraph("A Proposal and Design Report", ParagraphStyle("CoverType", fontName="Times-Italic", fontSize=12, leading=16, alignment=TA_CENTER)),
        Spacer(1, 0.55 * inch),
    ]
    cover_data = [
        ["Course", "[COURSE CODE AND COURSE TITLE]"],
        ["Faculty Guide", "[FACULTY GUIDE NAME AND DESIGNATION]"],
        ["Submitted by", "[STUDENT NAME(S) AND ROLL NUMBER(S)]"],
        ["Academic Year", "[2026-2027]"],
        ["Submission Date", "[DD MONTH YYYY]"],
    ]
    cover_table = Table(cover_data, colWidths=[1.35 * inch, 4.25 * inch])
    cover_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (0, -1), "Times-Bold"),
                ("FONTNAME", (1, 0), (1, -1), "Times-Roman"),
                ("FONTSIZE", (0, 0), (-1, -1), 10.5),
                ("GRID", (0, 0), (-1, -1), 0.75, colors.black),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    story.extend([cover_table, PageBreak(), Paragraph("Table of Contents", styles["Heading1"])])
    toc = TableOfContents()
    toc.levelStyles = [
        ParagraphStyle("TOC1", fontName="Times-Roman", fontSize=10.5, leading=14, leftIndent=0, firstLineIndent=0, spaceBefore=4),
    ]
    story.extend([toc, PageBreak()])

    table_text = ParagraphStyle("TableText", fontName="Times-Roman", fontSize=7.5, leading=9, alignment=TA_LEFT)
    table_head = ParagraphStyle("TableHead", fontName="Times-Bold", fontSize=8, leading=9, textColor=colors.white, alignment=TA_LEFT)

    for block in blocks:
        if block.kind == "heading1":
            story.append(Paragraph(_pdf_inline(str(block.value)), styles["Heading1"]))
        elif block.kind == "heading2":
            story.append(Paragraph(_pdf_inline(str(block.value)), styles["Heading2"]))
        elif block.kind == "paragraph":
            story.append(Paragraph(_pdf_inline(str(block.value)), styles["ReportBody"]))
        elif block.kind == "quote":
            story.append(Paragraph(_pdf_inline(str(block.value)), styles["Quote"]))
        elif block.kind == "list":
            marker, text = block.value
            story.append(Paragraph(f"{html.escape(marker)}&nbsp;&nbsp;{_pdf_inline(str(text))}", styles["List"]))
        elif block.kind == "code":
            _language, code = block.value
            story.append(
                KeepTogether(
                    [
                        Spacer(1, 4),
                        Preformatted(str(code), ParagraphStyle("Code", fontName="Courier", fontSize=7.5, leading=10, leftIndent=12)),
                        Spacer(1, 6),
                    ]
                )
            )
        elif block.kind == "image":
            caption, path = block.value
            image_path = HERE / str(path)
            with Image.open(image_path) as source_image:
                aspect = source_image.width / source_image.height
            max_width = 6.7 * inch
            max_height = 5.5 * inch
            if max_width / max_height > aspect:
                image_height = max_height
                image_width = image_height * aspect
            else:
                image_width = max_width
                image_height = image_width / aspect
            image = ReportImage(str(image_path), width=image_width, height=image_height)
            image.hAlign = "CENTER"
            caption_style = ParagraphStyle(
                "Caption", fontName="Times-Italic", fontSize=9, leading=11, alignment=TA_CENTER
            )
            story.extend([image, Spacer(1, 3), Paragraph(str(caption), caption_style), Spacer(1, 7)])
        elif block.kind == "table":
            rows = block.value
            col_count = len(rows[0])
            widths = [6.7 * inch / col_count] * col_count
            if col_count == 3:
                widths = [1.25 * inch, 2.65 * inch, 2.8 * inch]
            formatted = []
            for row_index, row in enumerate(rows):
                style = table_head if row_index == 0 else table_text
                formatted.append([Paragraph(_pdf_inline(cell), style) for cell in row])
            table = Table(formatted, colWidths=widths, repeatRows=1, hAlign="CENTER")
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.black),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            story.extend([table, Spacer(1, 7)])

    document = ReportTemplate(
        PDF_OUTPUT,
        pagesize=A4,
        rightMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.65 * inch,
        title="ShaderMesh Project Review 1",
        author="[STUDENT NAME(S)]",
    )
    document.multiBuild(story)


PPT_INK = PptxRGBColor(24, 24, 27)
PPT_MUTED = PptxRGBColor(82, 82, 91)
PPT_LINE = PptxRGBColor(161, 161, 170)
PPT_SOFT = PptxRGBColor(244, 244, 245)
PPT_AMBER = PptxRGBColor(180, 83, 9)
PPT_AMBER_SOFT = PptxRGBColor(255, 251, 235)


def _ppt_text_frame(
    shape,
    text: str,
    size: int,
    *,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    color=PPT_INK,
    font_name: str = "Arial",
):
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    for run in paragraph.runs:
        run.font.name = font_name
        run.font.size = PptxPt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return frame


def _ppt_text(slide, x, y, width, height, value, size, *, bold=False, color=PPT_INK, align=PP_ALIGN.LEFT, font_name="Arial"):
    shape = slide.shapes.add_textbox(PptxInches(x), PptxInches(y), PptxInches(width), PptxInches(height))
    frame = _ppt_text_frame(shape, value, size, bold=bold, align=align, color=color, font_name=font_name)
    frame.margin_left = PptxInches(0)
    frame.margin_right = PptxInches(0)
    frame.margin_top = PptxInches(0)
    frame.margin_bottom = PptxInches(0)
    return shape


def _ppt_rect(slide, x, y, width, height, *, fill=None, line=PPT_LINE, line_width=1.0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        PptxInches(x),
        PptxInches(y),
        PptxInches(width),
        PptxInches(height),
    )
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = PptxPt(line_width)
    shape.shadow.inherit = False
    return shape


def _ppt_base_slide(presentation, blank, title: str, index: int):
    slide = presentation.slides.add_slide(blank)
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = PptxRGBColor(255, 255, 255)
    _ppt_rect(slide, 0.62, 0.42, 0.07, 0.55, fill=PPT_AMBER, line=None)
    _ppt_text(slide, 0.86, 0.40, 11.7, 0.65, title, 26, bold=True)
    _ppt_rect(slide, 0.62, 1.13, 12.05, 0.015, fill=PPT_LINE, line=None)
    _ppt_text(slide, 0.65, 7.09, 11.7, 0.18, "SHADERMESH  |  PROJECT REVIEW 1", 8, color=PPT_MUTED)
    _ppt_text(slide, 12.15, 7.09, 0.5, 0.18, str(index), 8, color=PPT_MUTED, align=PP_ALIGN.RIGHT)
    return slide


def _ppt_picture_fit(slide, path: Path, x: float, y: float, width: float, height: float):
    with Image.open(path) as source:
        aspect = source.width / source.height
    box_aspect = width / height
    if aspect > box_aspect:
        rendered_width = width
        rendered_height = width / aspect
    else:
        rendered_height = height
        rendered_width = height * aspect
    left = x + (width - rendered_width) / 2
    top = y + (height - rendered_height) / 2
    slide.shapes.add_picture(
        str(path),
        PptxInches(left),
        PptxInches(top),
        width=PptxInches(rendered_width),
        height=PptxInches(rendered_height),
    )


def _ppt_row(slide, y: float, number: str, title: str, detail: str, *, x=1.0, width=11.3):
    _ppt_text(slide, x, y, 0.7, 0.42, number, 19, bold=True, color=PPT_AMBER)
    _ppt_text(slide, x + 0.82, y, 4.0, 0.42, title, 19, bold=True)
    _ppt_text(slide, x + 4.65, y, width - 4.65, 0.52, detail, 17, color=PPT_MUTED)
    _ppt_rect(slide, x + 0.82, y + 0.55, width - 0.82, 0.012, fill=PPT_LINE, line=None)


def _ppt_metric(slide, x: float, value: str, label: str):
    _ppt_text(slide, x, 5.55, 1.7, 0.55, value, 30, bold=True, color=PPT_AMBER, align=PP_ALIGN.CENTER)
    _ppt_text(slide, x, 6.15, 1.7, 0.48, label, 14, color=PPT_MUTED, align=PP_ALIGN.CENTER)


def build_pptx() -> None:
    presentation = Presentation()
    presentation.slide_width = PptxInches(13.333)
    presentation.slide_height = PptxInches(7.5)
    blank = presentation.slide_layouts[6]

    slide = presentation.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PptxRGBColor(255, 255, 255)
    _ppt_rect(slide, 1.0, 1.0, 0.08, 5.5, fill=PPT_AMBER, line=None)
    _ppt_text(slide, 1.45, 1.15, 10.7, 0.35, "PROJECT REVIEW 1", 16, bold=True, color=PPT_AMBER)
    _ppt_text(slide, 1.45, 1.7, 10.7, 0.9, "SHADERMESH", 42, bold=True)
    _ppt_text(slide, 1.45, 2.65, 10.4, 0.8, "A decentralized marketplace for renting compute", 25, color=PPT_MUTED)
    _ppt_text(slide, 1.45, 4.65, 10.4, 0.8, "[STUDENT NAME(S) AND ROLL NUMBER(S)]\n[INSTITUTION NAME]", 17, color=PPT_MUTED)
    _ppt_text(slide, 1.45, 6.35, 10.4, 0.3, "Proposal and Design", 14, bold=True, color=PPT_AMBER)

    slide = _ppt_base_slide(presentation, blank, "Idle hardware and compute demand lack a neutral agreement layer", 2)
    _ppt_text(slide, 0.9, 1.65, 3.2, 0.4, "UNDERUSED SUPPLY", 16, bold=True, color=PPT_MUTED)
    _ppt_text(slide, 0.9, 2.15, 3.2, 1.2, "Gaming PCs\nLaboratory workstations\nSmall GPU providers", 22, bold=True)
    _ppt_text(slide, 9.25, 1.65, 3.2, 0.4, "COMPUTE DEMAND", 16, bold=True, color=PPT_MUTED)
    _ppt_text(slide, 9.25, 2.15, 3.2, 1.2, "LLM inference\nRendering\nBatch workloads", 22, bold=True)
    _ppt_rect(slide, 4.6, 1.65, 4.1, 3.65, fill=PPT_AMBER_SOFT, line=PPT_AMBER, line_width=1.5)
    _ppt_text(slide, 5.0, 2.0, 3.3, 0.45, "THE TRUST GAP", 20, bold=True, color=PPT_AMBER, align=PP_ALIGN.CENTER)
    for row, item in enumerate(("Who holds payment?", "Was service available?", "Is the hardware genuine?", "What happens after failure?")):
        _ppt_text(slide, 5.05, 2.75 + row * 0.55, 3.2, 0.35, item, 18, align=PP_ALIGN.CENTER)
    _ppt_text(slide, 2.0, 5.75, 9.3, 0.55, "A public contract layer can make these rules visible and enforceable without running the workload on chain.", 21, bold=True, align=PP_ALIGN.CENTER)

    slide = _ppt_base_slide(presentation, blank, "Five objectives define a testable prototype", 3)
    objectives = [
        ("01", "Rental market", "Listings, escrow, interval settlement, refund, and termination"),
        ("02", "Provider trust", "Stake, attestation, reputation, and defined slashing paths"),
        ("03", "AI workloads", "OpenAI-compatible inference and accepted container execution"),
        ("04", "Verifiable tasks", "Commitments and disputes, then randomized verification hardening"),
        ("05", "Evaluation", "Tests, multi-laptop operation, gas, security, and economics"),
    ]
    for row, values in enumerate(objectives):
        _ppt_row(slide, 1.5 + row * 1.0, *values)

    slide = _ppt_base_slide(presentation, blank, "Prior systems solve different parts of the trust problem", 4)
    _ppt_text(slide, 0.85, 1.48, 2.25, 0.4, "CAPABILITY", 15, bold=True, color=PPT_MUTED)
    _ppt_text(slide, 3.25, 1.48, 3.4, 0.4, "REPRESENTATIVE WORK", 15, bold=True, color=PPT_MUTED)
    _ppt_text(slide, 6.85, 1.48, 5.55, 0.4, "LESSON FOR SHADERMESH", 15, bold=True, color=PPT_MUTED)
    literature = [
        ("Settlement", "Ethereum", "Keep consensus focused on value and state transitions"),
        ("Compute markets", "Golem, Akash, Vast.ai", "Resource matching does not settle every trust question"),
        ("Verification", "Truebit, Livepeer", "Re-execution applies only to reproducible work"),
        ("Attestation", "iExec", "Hardware evidence introduces its own trust assumptions"),
        ("GPU workloads", "Render Network", "Domain-specific clients trade generality for control"),
        ("Inference", "vLLM", "Efficient serving still needs market and settlement layers"),
    ]
    for row, values in enumerate(literature):
        y = 1.98 + row * 0.76
        _ppt_rect(slide, 0.82, y + 0.57, 11.65, 0.01, fill=PPT_LINE, line=None)
        _ppt_text(slide, 0.85, y, 2.2, 0.5, values[0], 17, bold=True)
        _ppt_text(slide, 3.25, y, 3.3, 0.5, values[1], 17, color=PPT_AMBER if row in (0, 2) else PPT_INK)
        _ppt_text(slide, 6.85, y, 5.5, 0.5, values[2], 16, color=PPT_MUTED)

    slide = presentation.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PptxRGBColor(255, 255, 255)
    _ppt_picture_fit(slide, SYSTEM_CONTEXT_OUTPUT, 0.35, 0.15, 12.63, 6.85)
    _ppt_text(slide, 0.65, 7.09, 11.7, 0.18, "SHADERMESH  |  PROJECT REVIEW 1", 8, color=PPT_MUTED)
    _ppt_text(slide, 12.15, 7.09, 0.5, 0.18, "5", 8, color=PPT_MUTED, align=PP_ALIGN.RIGHT)

    slide = _ppt_base_slide(presentation, blank, "ShaderMesh uses two workload-specific trust models", 6)
    columns = ((0.8, "TIME-BASED RENTAL", "Was service available?", "Timely availability evidence", "Escrow, refund, SLA slashing"), (6.75, "DETERMINISTIC TASK", "Does execution match?", "Commitment and re-execution", "Challenge, dispute, slashing"))
    for x, heading, question, evidence, outcome in columns:
        _ppt_text(slide, x, 1.55, 5.65, 0.45, heading, 21, bold=True, color=PPT_AMBER)
        for row, (label, value) in enumerate((("CORE QUESTION", question), ("EVIDENCE", evidence), ("ECONOMIC OUTCOME", outcome))):
            y = 2.25 + row * 1.15
            _ppt_text(slide, x, y, 5.5, 0.3, label, 14, bold=True, color=PPT_MUTED)
            _ppt_text(slide, x, y + 0.38, 5.5, 0.45, value, 19, bold=True)
            _ppt_rect(slide, x, y + 0.9, 5.45, 0.01, fill=PPT_LINE, line=None)
    _ppt_rect(slide, 0.8, 5.95, 11.6, 0.62, fill=PPT_SOFT, line=None)
    _ppt_text(slide, 1.05, 6.13, 11.1, 0.3, "Shared foundation: provider stake | public reputation | contract escrow | auditable outcomes", 18, bold=True, align=PP_ALIGN.CENTER)

    slide = presentation.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PptxRGBColor(255, 255, 255)
    _ppt_picture_fit(slide, RENTAL_SEQUENCE_OUTPUT, 0.35, 0.15, 12.63, 6.85)
    _ppt_text(slide, 0.65, 7.09, 11.7, 0.18, "SHADERMESH  |  PROJECT REVIEW 1", 8, color=PPT_MUTED)
    _ppt_text(slide, 12.15, 7.09, 0.5, 0.18, "7", 8, color=PPT_MUTED, align=PP_ALIGN.RIGHT)

    slide = _ppt_base_slide(presentation, blank, "Ethereum controls value while compute remains off chain", 8)
    layers = [
        (1.55, "PUBLIC CONSENSUS", "Ethereum proof of stake", PPT_SOFT, PPT_LINE),
        (2.75, "APPLICATION TRUST", "SmeshToken | ProviderRegistry | RentalMarket | JobMarket | Verification", PPT_AMBER_SOFT, PPT_AMBER),
        (3.95, "OFF-CHAIN COORDINATION", "Dashboard | Orchestrator | Provider agent | IPFS adapter", PPT_SOFT, PPT_LINE),
        (5.15, "COMPUTE EXECUTION", "Constrained Docker runtime | CPU / optional GPU | model server", PptxRGBColor(255, 255, 255), PPT_LINE),
    ]
    for y, label, detail, fill, border in layers:
        _ppt_rect(slide, 1.05, y, 11.2, 0.82, fill=fill, line=border, line_width=1.25)
        _ppt_text(slide, 1.35, y + 0.18, 2.8, 0.34, label, 15, bold=True, color=PPT_AMBER if border == PPT_AMBER else PPT_MUTED)
        _ppt_text(slide, 4.15, y + 0.16, 7.75, 0.4, detail, 19, bold=True)

    slide = presentation.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PptxRGBColor(255, 255, 255)
    _ppt_picture_fit(slide, ARCHITECTURE_OUTPUT, 0.35, 0.15, 12.63, 6.85)
    _ppt_text(slide, 0.65, 7.09, 11.7, 0.18, "SHADERMESH  |  PROJECT REVIEW 1", 8, color=PPT_MUTED)
    _ppt_text(slide, 12.15, 7.09, 0.5, 0.18, "9", 8, color=PPT_MUTED, align=PP_ALIGN.RIGHT)

    slide = presentation.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PptxRGBColor(255, 255, 255)
    _ppt_picture_fit(slide, TRUST_BOUNDARY_OUTPUT, 0.35, 0.15, 12.63, 6.85)
    _ppt_text(slide, 0.65, 7.09, 11.7, 0.18, "SHADERMESH  |  PROJECT REVIEW 1", 8, color=PPT_MUTED)
    _ppt_text(slide, 12.15, 7.09, 0.5, 0.18, "10", 8, color=PPT_MUTED, align=PP_ALIGN.RIGHT)

    slide = _ppt_base_slide(presentation, blank, "Implementation follows a contract-first sequence", 11)
    steps = (("1", "Trust model"), ("2", "State machines"), ("3", "Contract tests"), ("4", "Agent + runtime"), ("5", "Indexer + API"), ("6", "Dashboard"), ("7", "Sepolia"))
    for idx, (number, label) in enumerate(steps):
        x = 0.75 + idx * 1.75
        if idx < len(steps) - 1:
            _ppt_rect(slide, x + 0.55, 3.1, 1.25, 0.04, fill=PPT_LINE, line=None)
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, PptxInches(x), PptxInches(2.72), PptxInches(0.72), PptxInches(0.72))
        circle.fill.solid()
        circle.fill.fore_color.rgb = PPT_AMBER if idx < 6 else PptxRGBColor(255, 255, 255)
        circle.line.color.rgb = PPT_AMBER
        circle.shadow.inherit = False
        _ppt_text(slide, x, 2.88, 0.72, 0.25, number, 16, bold=True, color=PptxRGBColor(255, 255, 255) if idx < 6 else PPT_AMBER, align=PP_ALIGN.CENTER)
        _ppt_text(slide, x - 0.3, 3.65, 1.35, 0.75, label, 16, bold=True, align=PP_ALIGN.CENTER)
    _ppt_text(slide, 1.2, 5.25, 10.9, 0.55, "Each off-chain component was added only after its contract boundary and failure path were defined.", 21, bold=True, align=PP_ALIGN.CENTER)

    slide = _ppt_base_slide(presentation, blank, "The stack separates settlement, coordination, and execution", 12)
    stack = [
        ("INTERFACE", "Next.js | React | Tailwind | viem"),
        ("COORDINATION", "Node.js 20 | TypeScript | ethers"),
        ("PROVIDER", "Python 3.11 | web3.py | Docker SDK"),
        ("SETTLEMENT", "Solidity 0.8.24 | Hardhat | OpenZeppelin"),
        ("DATA + INFRA", "IPFS | Kubo | PostgreSQL | Docker Compose"),
    ]
    for row, (label, detail) in enumerate(stack):
        y = 1.48 + row * 1.0
        _ppt_rect(slide, 1.0, y, 2.25, 0.72, fill=PPT_AMBER_SOFT if label == "SETTLEMENT" else PPT_SOFT, line=PPT_AMBER if label == "SETTLEMENT" else PPT_LINE)
        _ppt_text(slide, 1.2, y + 0.2, 1.85, 0.3, label, 15, bold=True, color=PPT_AMBER if label == "SETTLEMENT" else PPT_MUTED, align=PP_ALIGN.CENTER)
        _ppt_text(slide, 3.65, y + 0.16, 8.35, 0.4, detail, 20, bold=True)
        _ppt_rect(slide, 3.65, y + 0.68, 8.35, 0.01, fill=PPT_LINE, line=None)

    slide = _ppt_base_slide(presentation, blank, "Operational inputs replace a training dataset", 13)
    _ppt_text(slide, 0.9, 1.5, 5.55, 0.42, "INPUTS", 18, bold=True, color=PPT_AMBER)
    _ppt_text(slide, 6.9, 1.5, 5.55, 0.42, "EVALUATION", 18, bold=True, color=PPT_AMBER)
    left_items = ("Lease duration and listing ID", "Resource specification and manifest", "Model ID, prompt, and sampling parameters", "Task CID and output commitment", "Simulator reward, bond, and spot-check rate")
    right_items = ("Escrow conservation and refunds", "Heartbeat and outage response", "Gas use by state transition", "Sandbox and GPU-access enforcement", "Route correctness, latency, and verification cost")
    for row, item in enumerate(left_items):
        _ppt_text(slide, 0.95, 2.15 + row * 0.75, 5.3, 0.45, f"{row + 1:02d}   {item}", 17, bold=row == 0)
    for row, item in enumerate(right_items):
        _ppt_text(slide, 6.95, 2.15 + row * 0.75, 5.3, 0.45, f"{row + 1:02d}   {item}", 17, bold=row == 0)
    _ppt_rect(slide, 0.9, 6.15, 11.55, 0.55, fill=PPT_SOFT, line=None)
    _ppt_text(slide, 1.05, 6.3, 11.2, 0.25, "Dataset applicability: not applicable. ShaderMesh is a marketplace protocol, not a trained model.", 16, bold=True, align=PP_ALIGN.CENTER)

    slide = _ppt_base_slide(presentation, blank, "The local prototype completes the main transaction path", 14)
    pipeline = (("DEPLOY", "contracts"), ("SEED", "providers"), ("LEASE", "escrow"), ("RUN", "container"), ("ROUTE", "inference"))
    for idx, (title, detail) in enumerate(pipeline):
        x = 0.85 + idx * 2.45
        _ppt_rect(slide, x, 2.0, 1.85, 1.05, fill=PPT_AMBER_SOFT if idx in (0, 2) else PPT_SOFT, line=PPT_AMBER if idx in (0, 2) else PPT_LINE)
        _ppt_text(slide, x, 2.2, 1.85, 0.3, title, 18, bold=True, color=PPT_AMBER if idx in (0, 2) else PPT_INK, align=PP_ALIGN.CENTER)
        _ppt_text(slide, x, 2.58, 1.85, 0.25, detail, 15, color=PPT_MUTED, align=PP_ALIGN.CENTER)
        if idx < len(pipeline) - 1:
            _ppt_rect(slide, x + 1.85, 2.52, 0.6, 0.04, fill=PPT_LINE, line=None)
    _ppt_metric(slide, 1.35, "7", "contract tests")
    _ppt_metric(slide, 3.65, "12", "agent tests")
    _ppt_metric(slide, 5.95, "17", "orchestrator tests")
    _ppt_metric(slide, 8.25, "3", "seeded listings")
    _ppt_metric(slide, 10.55, "1", "Compose command")

    slide = _ppt_base_slide(presentation, blank, "Current evidence proves integration, not production trust", 15)
    boundary_columns = [
        (0.8, "REAL", ("contract transactions", "escrow and settlement", "wallet calls and APIs", "container lifecycle"), PPT_INK),
        (4.55, "SEEDED", ("provider identities", "GPU declarations", "operator attestation", "echo inference"), PPT_MUTED),
        (8.3, "MUST BE HARDENED", ("timely heartbeat evidence", "random challenger selection", "commit-reveal and quorum", "real GPU benchmarks"), PPT_AMBER),
    ]
    for x, heading, items, color in boundary_columns:
        _ppt_text(slide, x, 1.55, 3.25, 0.4, heading, 17, bold=True, color=color)
        _ppt_rect(slide, x, 2.02, 3.25, 0.025, fill=color, line=None)
        for row, item in enumerate(items):
            _ppt_text(slide, x, 2.4 + row * 0.78, 3.25, 0.45, item, 17, bold=row == 0)
    _ppt_text(slide, 1.05, 6.05, 11.2, 0.55, "The documented limitations are part of the engineering result, not hidden assumptions.", 20, bold=True, align=PP_ALIGN.CENTER)

    slide = _ppt_base_slide(presentation, blank, "The next milestone is a reproducible multi-laptop demonstration", 16)
    outcomes = (("1", "Sepolia", "public contract deployment"), ("2", "NVIDIA nodes", "real hardware discovery"), ("3", "Accepted workload", "lease-bound container"), ("4", "Real model", "Ollama or vLLM"), ("5", "Failure test", "refund and slashing"))
    for idx, (number, title, detail) in enumerate(outcomes):
        x = 0.72 + idx * 2.5
        _ppt_text(slide, x, 1.75, 0.45, 0.35, number, 18, bold=True, color=PPT_AMBER)
        _ppt_rect(slide, x, 2.2, 2.0, 0.025, fill=PPT_LINE, line=None)
        _ppt_text(slide, x, 2.55, 2.05, 0.45, title, 19, bold=True)
        _ppt_text(slide, x, 3.15, 2.05, 0.75, detail, 16, color=PPT_MUTED)
    _ppt_text(slide, 1.1, 5.55, 11.1, 0.55, "Success means the full flow works across physical machines with explicit evidence and repeatable setup.", 20, bold=True, align=PP_ALIGN.CENTER)

    slide = _ppt_base_slide(presentation, blank, "The remaining work fits a four-week integration plan", 17)
    weeks = list(range(1, 13))
    for idx, week in enumerate(weeks):
        _ppt_text(slide, 3.2 + idx * 0.72, 1.45, 0.55, 0.3, str(week), 13, bold=True, color=PPT_MUTED, align=PP_ALIGN.CENTER)
    work = [
        ("Requirements + architecture", 1, 2, PPT_INK),
        ("Core contracts", 3, 4, PPT_INK),
        ("Agent + inference", 5, 6, PPT_INK),
        ("Orchestrator + dashboard", 7, 8, PPT_INK),
        ("Sepolia + multi-laptop", 9, 9, PPT_AMBER),
        ("GPU + model integration", 10, 10, PPT_LINE),
        ("Adversarial testing", 11, 11, PPT_LINE),
        ("Evaluation + delivery", 12, 12, PPT_LINE),
    ]
    for row, (label, start, end, color) in enumerate(work):
        y = 1.95 + row * 0.58
        _ppt_text(slide, 0.8, y, 2.25, 0.3, label, 14, bold=True if color in (PPT_INK, PPT_AMBER) else False, color=PPT_MUTED)
        _ppt_rect(slide, 3.2, y + 0.03, 8.47, 0.28, fill=PPT_SOFT, line=None)
        _ppt_rect(slide, 3.2 + (start - 1) * 0.72, y + 0.03, (end - start + 1) * 0.72 - 0.05, 0.28, fill=color, line=None)
    _ppt_text(slide, 3.2, 6.65, 8.5, 0.25, "Completed        Next milestone        Planned", 14, color=PPT_MUTED, align=PP_ALIGN.CENTER)

    slide = _ppt_base_slide(presentation, blank, "Selected sources establish the design baseline", 18)
    references = [
        "[1] Buterin, Ethereum Whitepaper, 2014",
        "[2] Golem official documentation",
        "[3] Teutsch and Reitwiessner, Truebit, 2019",
        "[4] Petkanics and Tang, Livepeer Whitepaper, 2017",
        "[5] iExec Proof-of-Contribution documentation",
        "[6] Akash Network documentation",
        "[7] Render Network Whitepaper, 2023",
        "[8] Protocol Labs, Filecoin, 2017",
        "[9] Kwon et al., vLLM, SOSP 2023",
        "[10] Vast.ai official documentation",
        "[11] OpenZeppelin Contracts documentation",
    ]
    for idx, reference in enumerate(references):
        column = 0 if idx < 6 else 1
        row = idx if column == 0 else idx - 6
        _ppt_text(slide, 0.9 + column * 6.05, 1.55 + row * 0.82, 5.55, 0.58, reference, 16, color=PPT_INK)
    _ppt_text(slide, 0.9, 6.55, 11.5, 0.25, "Full citations and access dates are included in the report.", 14, color=PPT_MUTED)

    slide = presentation.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PPT_INK
    _ppt_rect(slide, 0.95, 1.0, 0.08, 5.5, fill=PPT_AMBER, line=None)
    _ppt_text(slide, 1.4, 1.25, 10.8, 0.9, "Compute stays off chain.\nTrust and settlement stay auditable.", 34, bold=True, color=PptxRGBColor(255, 255, 255))
    _ppt_text(slide, 1.4, 3.25, 10.4, 1.25, "Rental mode: availability evidence + escrow\nTask mode: reproducible output + dispute", 22, color=PptxRGBColor(228, 228, 231))
    _ppt_text(slide, 1.4, 5.65, 10.4, 0.5, "Next: Sepolia + separate NVIDIA provider laptops", 18, bold=True, color=PptxRGBColor(251, 191, 36))
    _ppt_text(slide, 11.95, 7.09, 0.7, 0.18, "19", 8, color=PptxRGBColor(161, 161, 170), align=PP_ALIGN.RIGHT)

    presentation.core_properties.title = "ShaderMesh Project Review 1"
    presentation.core_properties.author = "[STUDENT NAME(S)]"
    presentation.save(PPTX_OUTPUT)


def main() -> None:
    blocks = parse_report()
    build_review_visuals()
    build_docx(blocks)
    build_pdf(blocks)
    build_pptx()
    print(DOCX_OUTPUT)
    print(PDF_OUTPUT)
    print(PPTX_OUTPUT)


if __name__ == "__main__":
    main()
