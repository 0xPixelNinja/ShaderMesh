"""Document helpers for the ShaderMesh final report (VIT-AP SCOPE format).

Adapted from the helper approach used for earlier VIT-AP reports. Provides fonts,
fields (TOC, list of figures and tables), captions numbered by chapter, tables,
section breaks with page-number restarts, and a two-pass label system so that
figures and tables can be referenced in the text before they appear.
"""
import re
import tempfile
import pathlib
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

FONT = "Times New Roman"
TEXT_W = 14.5  # cm: 21.0 - 3.5 - 3.0
TOK = re.compile(r"(\*\*.+?\*\*|`.+?`|__.+?__)")
URL = re.compile(r"(https?://\S+)")
CACHE = pathlib.Path(tempfile.mkdtemp(prefix="smesh_report_"))


def set_font(run, size=12, bold=None, italic=None, mono=False, underline=None):
    run.font.name = "Consolas" if mono else FONT
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:eastAsia"), run.font.name)
    run.font.size = Pt(size - 1.5 if mono else size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline


def fmt(p, align=AL.JUSTIFY, before=0, after=0, line=1.5, keep=False, left=None):
    f = p.paragraph_format
    p.alignment = align
    f.space_before, f.space_after, f.line_spacing = Pt(before), Pt(after), line
    if keep:
        f.keep_with_next = True
    if left is not None:
        f.left_indent = Cm(left)


def cell_shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def para_shade(p, fill="F2F2F2"):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def field(p, instr, result, size=12, bold=None):
    """Complex field with a cached result so it reads correctly before Word updates it."""
    def r_with(child):
        r = p.add_run()
        set_font(r, size, bold)
        r._r.append(child)
        return r
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    r_with(b)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r_with(it)
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    r_with(s)
    r = p.add_run(result)
    set_font(r, size, bold)
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    r_with(e)


def hyperlink(p, url, size=12):
    rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    f = OxmlElement("w:rFonts")
    f.set(qn("w:ascii"), FONT)
    f.set(qn("w:hAnsi"), FONT)
    rPr.append(f)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "0563C1")
    rPr.append(c)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = url
    r.append(t)
    h.append(r)
    p._p.append(h)


def set_section(sec, numfmt, start=None):
    sp = sec._sectPr
    for tag in ("w:pgNumType", "w:titlePg"):
        for e in sp.findall(qn(tag)):
            sp.remove(e)
    pn = OxmlElement("w:pgNumType")
    pn.set(qn("w:fmt"), numfmt)
    if start:
        pn.set(qn("w:start"), str(start))
    sp.find(qn("w:pgMar")).addnext(pn)
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin, sec.right_margin = Cm(3.5), Cm(3.0)
    sec.footer_distance = Cm(1.2)


def prepared_image(path, max_px=2000):
    """Downscale large PNGs so the DOCX stays small. Returns (path, aspect)."""
    src = pathlib.Path(path)
    im = Image.open(src)
    w, h = im.size
    out = CACHE / (src.parent.name + "_" + src.stem + ".png")
    if not out.exists():
        if w > max_px:
            im = im.resize((max_px, int(h * max_px / w)), Image.LANCZOS)
        im.convert("RGB").save(out, optimize=True)
    return out, w / h


def cropped_top(path, height_px, name):
    im = Image.open(path).convert("RGB")
    crop = im.crop((0, 0, im.size[0], min(height_px, im.size[1])))
    out = CACHE / name
    crop.save(out)
    return prepared_image(out)


class Report:
    def __init__(self, refs, labels=None):
        self.d = Document()
        self.refs, self.order = refs, []
        self.labels = labels if labels is not None else {}
        self.new_labels = {}
        self.referenced = set()
        self.chapter = 0
        self.sec = self.sub = 0
        self.fig_n = self.tab_n = 0
        self.strict = labels is not None
        st = self.d.styles
        n = st["Normal"]
        n.font.name = FONT
        n.font.size = Pt(12)
        n.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        n.paragraph_format.line_spacing = 1.5
        for nm in ("Figure Caption", "Table Caption"):
            s = st.add_style(nm, WD_STYLE_TYPE.PARAGRAPH)
            s.base_style = st["Normal"]
        for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12),
                           ("Figure Caption", 11), ("Table Caption", 11)):
            s = st[name]
            s.font.name = FONT
            s.font.size = Pt(size)
            s.font.bold = True
            s.font.italic = False
            s.font.color.rgb = RGBColor(0, 0, 0)
            rf = s.element.rPr.rFonts
            for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                if rf.get(qn("w:" + a)) is not None:
                    del rf.attrib[qn("w:" + a)]
            rf.set(qn("w:ascii"), FONT)
            rf.set(qn("w:hAnsi"), FONT)
            rf.set(qn("w:eastAsia"), FONT)
        h1 = st["Heading 1"].paragraph_format
        h1.alignment = AL.CENTER
        h1.space_before = Pt(0)
        h1.space_after = Pt(12)
        h1.page_break_before = True
        h1.keep_with_next = True
        h1.line_spacing = 1.5
        for nm in ("Heading 2", "Heading 3"):
            h = st[nm].paragraph_format
            h.space_before = Pt(10)
            h.space_after = Pt(4)
            h.keep_with_next = True
            h.line_spacing = 1.5
            h.alignment = AL.LEFT
        self.d.core_properties.title = "ShaderMesh: Final Project Report"
        self.d.core_properties.author = "ShaderMesh project team"
        set_section(self.d.sections[0], "decimal")

    # ---- text ----
    def cites(self, text):
        def rep(m):
            nums = []
            for k in m.group(1).split(","):
                k = k.strip().lstrip("@")
                if k not in self.refs:
                    raise KeyError(k)
                if k not in self.order:
                    self.order.append(k)
                nums.append(f"[{self.order.index(k) + 1}]")
            return ", ".join(nums)
        text = re.sub(r"\[(@[\w,@ ]+)\]", rep, text)

        def lab(m):
            key = m.group(2)
            self.referenced.add(key)
            if key not in self.labels:
                if self.strict:
                    raise KeyError("label " + key)
                return "??"
            lab_ = self.labels[key]
            return lab_
        return re.sub(r"\{([ft]):([\w-]+)\}", lab, text)

    def runs(self, p, text, size=12, bold=False):
        for part in TOK.split(self.cites(text)):
            if not part:
                continue
            if part.startswith("**"):
                set_font(p.add_run(part[2:-2]), size, True)
            elif part.startswith("`"):
                set_font(p.add_run(part[1:-1]), size, bold, mono=True)
            elif part.startswith("__"):
                set_font(p.add_run(part[2:-2]), size, bold, True)
            else:
                set_font(p.add_run(part), size, bold)

    def line(self, text, size=12, bold=False, italic=False, align=AL.CENTER, before=0, after=0,
             spacing=1.5, underline=None):
        p = self.d.add_paragraph()
        r = p.add_run(text)
        set_font(r, size, bold, italic, underline=underline)
        fmt(p, align, before, after, spacing)
        return p

    # headings with numbering
    def chapter_h1(self, title):
        self.chapter += 1
        self.sec = self.sub = 0
        self.fig_n = self.tab_n = 0
        return self.d.add_paragraph(f"CHAPTER {self.chapter}: {title}", style="Heading 1")

    def appendix_h1(self, letter, title):
        self.chapter = letter
        self.sec = self.sub = 0
        self.fig_n = self.tab_n = 0
        return self.d.add_paragraph(f"APPENDIX {letter}: {title}", style="Heading 1")

    def plain_h1(self, title):
        return self.d.add_paragraph(title, style="Heading 1")

    def h2(self, t):
        self.sec += 1
        self.sub = 0
        return self.d.add_paragraph(f"{self.chapter}.{self.sec} {t}", style="Heading 2")

    def h3(self, t):
        self.sub += 1
        return self.d.add_paragraph(f"{self.chapter}.{self.sec}.{self.sub} {t}", style="Heading 3")

    def p(self, t, line=1.5, after=6, size=12, keep=False):
        p = self.d.add_paragraph()
        self.runs(p, t, size)
        fmt(p, after=after, line=line, keep=keep)
        return p

    def bullets(self, items):
        for t in items:
            p = self.d.add_paragraph()
            self.runs(p, "•\t" + t)
            fmt(p, after=3, left=0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))

    def numbered(self, items):
        for i, t in enumerate(items, 1):
            p = self.d.add_paragraph()
            self.runs(p, f"{i}.\t" + t)
            fmt(p, after=3, left=0.9)
            p.paragraph_format.first_line_indent = Cm(-0.7)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))

    def eq(self, text):
        p = self.d.add_paragraph()
        r = p.add_run(text)
        set_font(r, 12, italic=True)
        fmt(p, AL.CENTER, 2, 8, 1.3)

    def code(self, lines, size=10.5):
        for i, t in enumerate(lines):
            p = self.d.add_paragraph()
            r = p.add_run(t)
            set_font(r, size + 1.5, mono=True)
            fmt(p, AL.LEFT, 0, 0 if i < len(lines) - 1 else 8, 1.0, left=0.2)
            para_shade(p)
            p.paragraph_format.keep_together = True
            if i < len(lines) - 1:
                p.paragraph_format.keep_with_next = True

    # captions
    def _label(self, kind, key):
        if kind == "Figure":
            self.fig_n += 1
            lab = f"Fig. {self.chapter}.{self.fig_n}"
        else:
            self.tab_n += 1
            lab = f"Table {self.chapter}.{self.tab_n}"
        self.new_labels[key] = lab
        if self.strict and key not in self.referenced:
            raise RuntimeError(f"{kind} '{key}' is not referenced in the text before it appears")
        return lab

    def caption(self, lab, text, above):
        p = self.d.add_paragraph(style="Table Caption" if above else "Figure Caption")
        shown = lab.upper() if above else lab
        r = p.add_run(f"{shown}: {self.cites(text)}")
        set_font(r, 11, True)
        fmt(p, AL.CENTER, 6 if above else 2, 4 if above else 10, 1.15, keep=above)
        return p

    def figure(self, key, path, cap, max_w=TEXT_W, max_h=12.0, prepared=None):
        lab = self._label("Figure", key)
        img, aspect = prepared if prepared else prepared_image(path)
        w = min(max_w, max_h * aspect)
        p = self.d.add_paragraph()
        p.add_run().add_picture(str(img), width=Cm(w))
        fmt(p, AL.CENTER, 6, 0, 1.0, keep=True)
        self.caption(lab, cap, above=False)

    def table(self, key, title, header, rows, widths, size=10.5, align_right=()):
        lab = self._label("Table", key)
        self.caption(lab, title, above=True)
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            self.runs(c.paragraphs[0], h, size, True)
            cell_shade(c, "D9D9D9")
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                self.runs(cells[i].paragraphs[0], v, size)
        scale = TEXT_W / sum(widths)
        n = len(t.rows)
        for ri, r in enumerate(t.rows):
            trPr = r._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            if ri == 0:
                trPr.append(OxmlElement("w:tblHeader"))
            for i, w in enumerate(widths):
                r.cells[i].width = Cm(w * scale)
            for c in r.cells:
                for para in c.paragraphs:
                    fmt(para, AL.LEFT, 2, 2, 1.0)
                    # keep the header and first rows with the caption; short tables stay whole
                    if ri < 2 or (n <= 12 and ri < n - 1):
                        para.paragraph_format.keep_with_next = True
            for i in align_right:
                for para in r.cells[i].paragraphs:
                    para.alignment = AL.RIGHT
        sp = self.d.add_paragraph()
        fmt(sp, after=0, line=1.0)
        sp.paragraph_format.space_after = Pt(6)

    # ---- structure ----
    def page_break(self):
        self.d.add_page_break()

    def new_section(self, numfmt, start=None):
        sec = self.d.add_section(WD_SECTION.NEW_PAGE)
        set_section(sec, numfmt, start)
        return sec

    def footer_numbers(self, sec):
        sec.footer.is_linked_to_previous = False
        p = sec.footer.paragraphs[0]
        p.alignment = AL.CENTER
        field(p, "PAGE", "1", 12)

    def references(self):
        self.plain_h1("REFERENCES")
        for i, k in enumerate(self.order, 1):
            p = self.d.add_paragraph()
            set_font(p.add_run(f"[{i}]\t"), 12)
            for part in URL.split(self.refs[k]):
                if not part:
                    continue
                if URL.fullmatch(part):
                    hyperlink(p, part)
                else:
                    self.runs(p, part, 12)
            fmt(p, AL.LEFT, 0, 5, 1.15, left=1.1)
            p.paragraph_format.first_line_indent = Cm(-1.1)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(1.1))
