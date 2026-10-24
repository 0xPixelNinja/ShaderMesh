#!/usr/bin/env python3
"""Build the ShaderMesh final project report (DOCX and PDF).

Run from anywhere:

    python docs/_scripts/build_report.py            # build DOCX, update fields, export PDF
    python docs/_scripts/build_report.py --no-pdf   # build DOCX only (no Word needed)

Output: docs/ShaderMesh-Final-Report.docx and docs/ShaderMesh-Final-Report.pdf

Format: VIT-AP SCOPE report guideline (A4, left 3.5 cm, right 3.0 cm, top and
bottom 2.54 cm, Times New Roman 12 pt, 1.5 spacing, justified). Front matter is
unnumbered or roman, arabic page numbers start at Chapter 1. Figures are
numbered "Fig. N.M" with the caption below, tables "TABLE N.M" with the title
above. The table of contents and the lists of tables and figures are Word
fields. Microsoft Word (through pywin32) updates them and exports the PDF.

Supporting modules in this folder: report_lib.py (helpers), report_refs.py
(references), report_content_a.py, report_content_b.py, report_content_c.py
(chapters). Figures come from docs/_report/figures (run report_figures.py
first if they are missing). Data comes from docs/_report/data.
"""
import sys
import pathlib

HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.enum.table import WD_TABLE_ALIGNMENT

from report_lib import Report, set_font, fmt, field, cell_shade, cropped_top
from report_refs import REFS
import report_content_a as A
import report_content_b as B
import report_content_c as C

# --------------------------------------------------------------------------
# Fill these in one place. Bracketed values are placeholders.
# --------------------------------------------------------------------------
MEMBERS = [
    ("Rakesh Kumar Rabha", "24BCC7002"),
    ("[MEMBER 2 NAME]", "[MEMBER 2 REG. NO.]"),
    ("[MEMBER 3 NAME]", "[MEMBER 3 REG. NO.]"),
]
GUIDE_NAME = "[GUIDE NAME]"
GUIDE_DESIGNATION = "[GUIDE DESIGNATION]"
HOD_NAME = "[HOD NAME]"

TITLE = "SHADERMESH: A DECENTRALIZED MARKETPLACE FOR RENTING GPU AND LLM COMPUTE ON ETHEREUM"
TITLE_CASE = "ShaderMesh: A Decentralized Marketplace for Renting GPU and LLM Compute on Ethereum"
MONTH_YEAR = "OCTOBER 2026"
PLACE = "Amaravati"
# --------------------------------------------------------------------------

ROOT = HERE.parents[1]
DOCS = ROOT / "docs"
FIG = DOCS / "_report" / "figures"
SCR = DOCS / "_report" / "screens"
DATA = DOCS / "_report" / "data"
LOGO = DOCS / "_report" / "assets" / "vitap-logo.jpeg"
OUT_DOCX = DOCS / "ShaderMesh-Final-Report.docx"
OUT_PDF = DOCS / "ShaderMesh-Final-Report.pdf"

ADDR = [
    ("SmeshToken", "0x4d08515811A8b14A69aE3049278EDB04e83645F1"),
    ("ProviderRegistry", "0xd723DFd9E63632e9533ea1946b57AdB6d00cfC3f"),
    ("JobMarket", "0x9F3B0526A9c665F1Df423145eaACCe5b2479Ee57"),
    ("Verification", "0x994927B22613bCB8c5A25fA6386E386F250D365F"),
    ("RentalMarket", "0x0CcaB81947Ee0Dc270908FE0cd99f280596f3Ac3"),
    ("InferenceMarket", "0x108fF3eaeB81592A04212F52bEe50F46a0abf9D1"),
]

ABBR = [
    ("API", "Application Programming Interface"), ("CI", "Continuous Integration"),
    ("CID", "Content Identifier (IPFS)"), ("EIP", "Ethereum Improvement Proposal"),
    ("EMA", "Exponential Moving Average"), ("ERC", "Ethereum Request for Comments (token standard)"),
    ("ETH", "Ether, the native currency of Ethereum"), ("EVM", "Ethereum Virtual Machine"),
    ("GPU", "Graphics Processing Unit"), ("IPFS", "InterPlanetary File System"),
    ("JSON-RPC", "JavaScript Object Notation Remote Procedure Call"), ("LLM", "Large Language Model"),
    ("RPC", "Remote Procedure Call"), ("SHA-256", "Secure Hash Algorithm, 256 bit"),
    ("SLA", "Service Level Agreement"), ("SMESH", "ShaderMesh utility token"),
    ("SSE", "Server-Sent Events"), ("SSIM", "Structural Similarity Index"),
    ("TEE", "Trusted Execution Environment"), ("WCAG", "Web Content Accessibility Guidelines"),
]

ABSTRACT = [
    "Demand for GPU compute has grown with the spread of large language models, while many capable GPUs sit idle. Renting capacity between strangers needs payment security, evidence of service, and enforcement of penalties, and these are normally supplied by a trusted marketplace operator. This project builds ShaderMesh, a decentralized marketplace for renting GPU and CPU compute on Ethereum, in which the blockchain is the trust, coordination, and payment layer and all compute runs off chain.",
    "ShaderMesh uses two explicit trust models over one staking, reputation, and slashing layer. Rentals cannot be re-executed and checked, so they are secured by streamed escrow, bonded provider stake, strict per-interval signed heartbeats, hardware attestation, and seeded benchmark challenges that slash free stake on failure. Deterministic jobs, which can be recomputed, use an optimistic verification game with spot checks and slashing. Managed LLM inference is sold through prepaid credit and an OpenAI-compatible gateway with signed requests.",
    "The system consists of six Solidity contracts, a Python provider agent, an attester service, a TypeScript orchestrator with a rebuildable Postgres projection, a Next.js web application, and a Monte Carlo simulator. It runs locally with Docker Compose and is deployed and source-verified on the Sepolia test network. The automated suites contain 42 contract tests, 58 agent tests, 54 orchestrator tests, 15 simulator tests, and 24 browser checks that include wallet writes.",
    "The evaluation reports gas for 27 core transactions, a simulation that reproduces the break-even spot-check rate of 0.5 and confirms conservation of escrow in rentals, and a live run on an RTX 4070 Laptop GPU provider that registered, staked, served a real model, and passed a live benchmark challenge with a score of 28,919. The simulation also shows that detection falls below the nominal spot-check rate in a mixed challenger pool, and that strict heartbeat termination penalizes brief outages. The report states the limits of the design: the attester is a trusted judge, LLM output is not verified on chain, and a heartbeat proves the liveness of a key and not the quality of service.",
]


def front_matter(R, d):
    # ---- cover ----
    R.line(TITLE, 18, True, before=0, after=14)
    R.line("A Project Report", 14, after=8)
    R.line("submitted by", 14, italic=True, after=4)
    for n, reg in MEMBERS:
        R.line(n.upper(), 14, True, after=0)
        R.line(f"({reg})", 12, True, after=4)
    R.line("in partial fulfilment of the requirements for the degree of", 14, italic=True, before=6, after=4)
    R.line("BACHELOR OF TECHNOLOGY", 16, True)
    R.line("in", 14, True)
    R.line("COMPUTER SCIENCE AND ENGINEERING", 14, True, after=4)
    R.line(f"Under the guidance of {GUIDE_NAME}, {GUIDE_DESIGNATION}", 12, italic=True, after=8, spacing=1.15)
    if LOGO.exists():
        lp = d.add_paragraph()
        lp.add_run().add_picture(str(LOGO), width=Cm(7))
        fmt(lp, AL.CENTER, 4, 8, 1.0)
    R.line("SCHOOL OF COMPUTER SCIENCE AND ENGINEERING", 14, True, after=2, spacing=1.15)
    R.line("VIT-AP UNIVERSITY, AMARAVATI, ANDHRA PRADESH", 12, True, after=8, spacing=1.15)
    R.line(MONTH_YEAR, 14)
    R.page_break()

    # ---- declaration ----
    R.line("DECLARATION BY THE CANDIDATES", 14, True, before=6, after=14)
    p = d.add_paragraph()
    R.runs(p, f"We hereby declare that the project report entitled “**{TITLE_CASE}**” submitted by us to VIT-AP University, Amaravati in partial fulfilment of the requirement for the award of the degree of **Bachelor of Technology** in **Computer Science and Engineering** is a record of bonafide project work carried out by us under the supervision of **{GUIDE_NAME}, {GUIDE_DESIGNATION}**. We further declare that the work reported in this report has not been submitted and will not be submitted, either in part or in full, for the award of any other degree or diploma in this institute or any other institute or university.", 13)
    fmt(p, AL.JUSTIFY, 0, 24, 1.5)
    R.line(f"Place: {PLACE}", 13, align=AL.LEFT)
    R.line("Date:", 13, align=AL.LEFT, after=18)
    for n, reg in MEMBERS:
        R.line("Signature of the student", 12, align=AL.RIGHT, spacing=1.0, before=10)
        R.line(f"Name: {n}", 12, align=AL.RIGHT, spacing=1.0)
        R.line(f"Reg. Number: {reg}", 12, align=AL.RIGHT, spacing=1.0)
    R.page_break()

    # ---- certificate ----
    R.line("School of Computer Science and Engineering", 14, True, before=0, after=6)
    R.line("BONAFIDE CERTIFICATE", 16, True, after=12, underline=True)
    p = d.add_paragraph()
    names = "; ".join(f"{n} ({reg})" for n, reg in MEMBERS)
    R.runs(p, f"This is to certify that the project report entitled “**{TITLE_CASE}**” submitted by {names} to VIT-AP University, Amaravati in partial fulfilment of the requirement for the award of the degree of **Bachelor of Technology** in **Computer Science and Engineering** is a record of bonafide project work undertaken by them under my supervision. The work fulfils the requirements as per the regulations of this Institute and, in my opinion, meets the necessary standards for submission. The contents of this report have not been submitted and will not be submitted either in part or in full, for the award of any other degree or diploma in this institute or any other institute or university.", 13)
    fmt(p, AL.JUSTIFY, 0, 30, 1.5)
    t = d.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = t.rows[0].cells
    for cell, lines in ((left, ["Signature of the Guide", f"Name: {GUIDE_NAME}", "GUIDE", GUIDE_DESIGNATION, "Date:"]),
                        (right, ["Signature of the Head of the Department", f"Name: {HOD_NAME}", "HEAD OF THE DEPARTMENT", "School of Computer Science and Engineering", "Date:"])):
        cell.text = ""
        for i, ln in enumerate(lines):
            para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            set_font(para.add_run(ln), 12, bold=(ln in ("GUIDE", "HEAD OF THE DEPARTMENT")))
            fmt(para, AL.LEFT, 0, 0, 1.15)
        cell.paragraphs[0].paragraph_format.space_before = Pt(36)
    for c in (left, right):
        c.width = Cm(7.25)
    sp = d.add_paragraph()
    fmt(sp, after=0, line=1.0)
    sp.paragraph_format.space_before = Pt(40)
    t2 = d.add_table(rows=1, cols=2)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    a, b = t2.rows[0].cells
    for cell, lines in ((a, ["Signature", "Internal Examiner(s)", "Date:"]), (b, ["Signature", "External Examiner(s)", "Date:"])):
        cell.text = ""
        for i, ln in enumerate(lines):
            para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            set_font(para.add_run(ln), 12, bold=(i == 1))
            fmt(para, AL.LEFT, 0, 0, 1.15)
        cell.paragraphs[0].paragraph_format.space_before = Pt(36)
        cell.width = Cm(7.25)

    # ---- section B: acknowledgement, abstract, lists (roman) ----
    secB = R.new_section("lowerRoman")
    R.footer_numbers(secB)
    R.plain_h1("ACKNOWLEDGEMENT")
    R.p(f"We thank our guide, {GUIDE_NAME}, {GUIDE_DESIGNATION}, School of Computer Science and Engineering, VIT-AP University, for guidance and review during the project, and for the feedback at Review 1 that shaped the move from a rendering marketplace to a compute-rental design. We thank the Head of the Department, {HOD_NAME}, and the faculty of the School of Computer Science and Engineering for the academic environment in which this work was done. We also acknowledge the maintainers of the open-source projects on which the system is built, including OpenZeppelin, Hardhat, Next.js, Ollama, IPFS, and PostgreSQL, whose documentation made it possible to build and test the system in the time available.", after=18)
    R.line(f"Place: {PLACE}", 12, align=AL.LEFT)
    R.line("Date:", 12, align=AL.LEFT, after=18)
    for n, reg in MEMBERS:
        R.line(f"{n} ({reg})", 12, align=AL.RIGHT, spacing=1.15)

    R.plain_h1("ABSTRACT")
    for s in ABSTRACT:
        R.p(s, line=2.0, after=0)

    def toc_page(title, instr, hint):
        p = d.add_paragraph()
        set_font(p.add_run(title), 14, True)
        fmt(p, AL.CENTER, 0, 12, 1.5)
        p.paragraph_format.page_break_before = True
        q = d.add_paragraph()
        field(q, instr, hint)
        fmt(q, AL.LEFT, 0, 0, 1.5)

    toc_page("TABLE OF CONTENTS", 'TOC \\o "1-3" \\h \\z \\u', "Update this field to build the table of contents.")
    R.plain_h1("LIST OF TABLES")
    q = d.add_paragraph()
    field(q, 'TOC \\h \\z \\t "Table Caption,1"', "Update this field to build the list of tables.")
    fmt(q, AL.LEFT, 0, 0, 1.5)
    R.plain_h1("LIST OF FIGURES")
    q = d.add_paragraph()
    field(q, 'TOC \\h \\z \\t "Figure Caption,1"', "Update this field to build the list of figures.")
    fmt(q, AL.LEFT, 0, 0, 1.5)

    R.plain_h1("LIST OF SYMBOLS, ABBREVIATIONS AND NOMENCLATURE")
    R.line("Abbreviations", 12, True, align=AL.LEFT, before=4, after=4, spacing=1.15)
    t = d.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, h in zip(t.rows[0].cells, ("Abbreviation", "Expansion")):
        c.text = ""
        set_font(c.paragraphs[0].add_run(h), 11, True)
        cell_shade(c, "D9D9D9")
    for a_, b_ in ABBR:
        cs = t.add_row().cells
        for c, v in zip(cs, (a_, b_)):
            c.text = ""
            set_font(c.paragraphs[0].add_run(v), 11)
    for r in t.rows:
        r.cells[0].width, r.cells[1].width = Cm(3.4), Cm(11.1)
        for c in r.cells:
            fmt(c.paragraphs[0], AL.LEFT, 1, 1, 1.0)
    sp = d.add_paragraph()
    fmt(sp, after=4, line=1.0)
    R.line("Symbols and Nomenclature", 12, True, align=AL.LEFT, before=4, after=4, spacing=1.15)
    R.bullets([
        "**R, B:** reward per task and bond per task, in SMESH.",
        "**s, s*:** spot-check probability and its break-even value, s* = R / (R + B).",
        "**V, D, h, h*:** value captured per interval by a false claim, bond at risk, detection probability, and break-even detection h* = V / (V + D).",
        "**Reputation:** integer from 0 to 10,000 updated by an exponential moving average with K = 5.",
        "**Provider, requester, attester:** the roles defined in Chapter 5.",
    ])


def build(labels):
    R = Report(REFS, labels)
    d = R.d
    front_matter(R, d)
    # section C: chapters, arabic numbering from 1
    secC = R.new_section("decimal", 1)
    S = {
        "landing": SCR / "landing.png",
        "landing-dark": SCR / "landing-dark.png",
        "landing_crop": cropped_top(SCR / "landing.png", 1800, "landing_top.png"),
        "landing_dark_crop": cropped_top(SCR / "landing-dark.png", 1800, "landing_dark_top.png"),
    }
    for k in ("dashboard", "rent", "models", "providers", "explorer-leases", "explorer-challenges",
              "explorer-events", "explorer-challenges-dark"):
        S[k] = SCR / f"{k}.png"
    A.chapter1(R, FIG)
    A.chapter2(R, FIG)
    A.chapter3(R, FIG)
    A.chapter4(R, FIG)
    A.chapter5(R, FIG)
    B.chapter6(R, FIG)
    B.chapter7(R, FIG)
    B.chapter8(R, FIG, S)
    B.chapter9(R, FIG)
    C.chapter10(R, FIG, S, DATA)
    C.chapter11(R, FIG)
    C.chapter12(R, FIG)
    R.references()
    C.appendices(R, FIG, ADDR)
    return R


def export_pdf(docx_path, pdf_path):
    import win32com.client
    w = win32com.client.DispatchEx("Word.Application")
    w.Visible = False
    w.DisplayAlerts = 0
    doc = w.Documents.Open(str(docx_path))
    try:
        for _ in range(3):
            doc.Repaginate()
            for t in doc.TablesOfContents:
                t.Update()
            for t in doc.TablesOfFigures:
                t.Update()
            doc.Fields.Update()
        doc.Save()
        doc.ExportAsFixedFormat(str(pdf_path), 17)
        return doc.ComputeStatistics(2)
    finally:
        doc.Close(0)
        w.Quit()


def main():
    no_pdf = "--no-pdf" in sys.argv
    first = build(None)          # pass 1 collects figure and table numbers
    final = build(first.new_labels)  # pass 2 resolves references and checks order
    final.d.save(str(OUT_DOCX))
    print("saved", OUT_DOCX, "figures/tables:", len(final.new_labels))
    if not no_pdf:
        pages = export_pdf(OUT_DOCX, OUT_PDF)
        print("saved", OUT_PDF, "total pages", pages)


if __name__ == "__main__":
    main()
