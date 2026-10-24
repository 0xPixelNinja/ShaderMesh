#!/usr/bin/env python3
"""Build the short ShaderMesh project report (proposal and design format).

Run from anywhere:

    python docs/_scripts/build_project_report.py            # DOCX and PDF
    python docs/_scripts/build_project_report.py --no-pdf   # DOCX only

Output: docs/ShaderMesh-Project-Report.docx and .pdf

The report follows the sixteen headings of the course project template, in
order. It reuses the helpers in report_lib.py, the reference list in
report_refs.py, the figures in docs/_report/figures (report_figures.py), and the
screenshots in docs/_report/screens. Team details are shared with
build_report.py so they are filled in one place.
"""
import sys
import pathlib

HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.enum.table import WD_TABLE_ALIGNMENT

from report_lib import Report, set_font, fmt, field, cell_shade
from report_refs import REFS
from build_report import MEMBERS, GUIDE_NAME, GUIDE_DESIGNATION, TITLE, MONTH_YEAR, ADDR, LOGO, export_pdf

ROOT = HERE.parents[1]
DOCS = ROOT / "docs"
FIG = DOCS / "_report" / "figures"
SCR = DOCS / "_report" / "screens"
OUT_DOCX = DOCS / "ShaderMesh-Project-Report.docx"
OUT_PDF = DOCS / "ShaderMesh-Project-Report.pdf"


def section(R, title):
    R.chapter += 1
    R.sec = R.sub = 0
    R.fig_n = R.tab_n = 0
    return R.d.add_paragraph(f"{R.chapter}. {title}", style="Heading 1")


def cover(R):
    d = R.d
    R.line(TITLE, 18, True, after=14)
    R.line("Project Report: Proposal and Design", 14, after=18)
    R.line("Submitted by", 13, italic=True, after=6)
    t = d.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, h in zip(t.rows[0].cells, ("Name", "Roll number")):
        c.text = ""
        set_font(c.paragraphs[0].add_run(h), 12, True)
        cell_shade(c, "D9D9D9")
    for n, reg in MEMBERS:
        cs = t.add_row().cells
        for c, v in zip(cs, (n, reg)):
            c.text = ""
            set_font(c.paragraphs[0].add_run(v), 12)
    for r in t.rows:
        r.cells[0].width, r.cells[1].width = Cm(7.0), Cm(6.0)
        for c in r.cells:
            fmt(c.paragraphs[0], AL.LEFT, 2, 2, 1.15)
    R.line(f"Under the guidance of {GUIDE_NAME}, {GUIDE_DESIGNATION}", 12, italic=True, before=18, after=12, spacing=1.15)
    if LOGO.exists():
        lp = d.add_paragraph()
        lp.add_run().add_picture(str(LOGO), width=Cm(6))
        fmt(lp, AL.CENTER, 6, 10, 1.0)
    R.line("SCHOOL OF COMPUTER SCIENCE AND ENGINEERING", 14, True, after=2, spacing=1.15)
    R.line("VIT-AP UNIVERSITY, AMARAVATI, ANDHRA PRADESH", 12, True, after=10, spacing=1.15)
    R.line(MONTH_YEAR, 14)

    p = d.add_paragraph()
    set_font(p.add_run("CONTENTS"), 14, True)
    fmt(p, AL.CENTER, 0, 12, 1.5)
    p.paragraph_format.page_break_before = True
    q = d.add_paragraph()
    field(q, 'TOC \\o "1-1" \\h \\z \\u', "Update this field to build the table of contents.")
    fmt(q, AL.LEFT, 0, 0, 1.5)


def body(R):
    F, S = FIG, SCR

    first = section(R, "PROJECT TITLE")
    first.paragraph_format.page_break_before = True
    R.p("**ShaderMesh: A Decentralized Marketplace for Renting GPU and LLM Compute on Ethereum**")
    R.p("ShaderMesh lets people with idle GPUs rent them to people who need compute, without a company in the middle holding the money. Ethereum smart contracts hold the escrow and the providers' collateral, pay providers one interval at a time, and penalize providers that stop serving. The workloads themselves, such as GPU containers and LLM inference, run off chain on the provider's machine.")

    section(R, "TEAM MEMBERS AND ROLL NUMBERS")
    R.p("{t:team} lists the team members.")
    R.table("team", "Team members", ["S. No.", "Name", "Roll number"],
            [[str(i), n, reg] for i, (n, reg) in enumerate(MEMBERS, 1)], [2, 8, 5])
    R.p(f"Guide: {GUIDE_NAME}, {GUIDE_DESIGNATION}, School of Computer Science and Engineering, VIT-AP University.")

    section(R, "PROBLEM STATEMENT")
    R.h2("What problem is being addressed")
    R.p("Demand for GPU compute has grown quickly with large language models, and cloud GPUs are expensive and often unavailable. At the same time many capable GPUs, in gaming PCs, laptops, and small labs, sit idle most of the day. Renting that capacity between strangers is hard for three reasons:")
    R.bullets([
        "**Payment risk.** The renter does not want to pay before the hours are served, and the owner does not want to serve before being paid.",
        "**No evidence of service.** Nothing records whether the GPU was actually available for each hour that was billed.",
        "**No enforcement.** If a provider goes offline or misrepresents its hardware, the renter has no recourse except a central platform's support desk.",
    ])
    R.p("Existing GPU marketplaces solve this by acting as the trusted operator: they hold funds, keep reputation scores, and decide disputes. {f:problem} contrasts that model with the approach taken in this project.")
    R.figure("problem", F / "problem.png", "A central operator compared with contracts as the trusted party", max_h=7.0)
    R.h2("Why blockchain and cryptocurrency are relevant")
    R.bullets([
        "**Escrow without a custodian.** A smart contract can hold the renter's deposit and release it per interval under fixed rules that neither party can change.",
        "**Collateral that can be slashed.** Providers lock tokens as stake. A contract can take part of that stake when a provider fails, which gives the provider a reason to stay honest.",
        "**A shared, public record.** Leases, heartbeats, payouts, challenges, and slashes are events on a public chain, so reputation is computed from history that anyone can check.",
        "**Programmable payments.** A token lets the contract pay small amounts every interval and refund the unused remainder automatically.",
    ])
    R.p("The blockchain is used only where shared economic state and deterministic enforcement are needed. Compute stays off chain, because executing GPU workloads on chain is neither possible nor useful.")

    section(R, "OBJECTIVES")
    R.numbered([
        "Design and implement smart contracts that hold rental escrow and provider stake, stream payment to providers per verified interval, refund unused time, and slash providers that breach the service terms.",
        "Build a trust model for compute that cannot be re-executed, based on signed per-interval heartbeats, hardware attestation, and seeded benchmark challenges that are judged and enforced on chain.",
        "Deliver real GPU and LLM compute through the marketplace: GPU containers run by a provider agent with Docker, and an OpenAI-compatible inference endpoint paid through prepaid credit.",
        "Provide a web application in which requesters rent capacity and providers stake and list hardware, with every value-moving action signed by the user's own wallet.",
        "Deploy the system to the Ethereum Sepolia test network, run it with a real GPU provider, and evaluate it with automated tests, gas measurements, and an economic simulation.",
    ])

    section(R, "MOTIVATION AND NEED FOR THE PROJECT")
    R.bullets([
        "**Cost and access.** Students, researchers, and small teams need GPUs for model inference and fine-tuning but cannot always pay cloud prices. Idle consumer GPUs can serve many of these workloads.",
        "**Trust between strangers.** A peer-to-peer market only works if both sides are protected. Escrow and stake replace the trust normally placed in a platform company.",
        "**A gap in existing systems.** Most decentralized compute projects either verify results by re-execution, which does not work for interactive rentals, or rely on a central operator for disputes. A rental market that is secured by availability evidence and slashing is less explored.",
        "**Learning value.** The project covers the full stack of a blockchain application: contract design and security, off-chain services, wallet-based user interfaces, testnet deployment, and economic analysis.",
    ])

    section(R, "LITERATURE SURVEY")
    R.p("The survey covers decentralized compute markets, verification protocols, and centralized GPU platforms. {t:lit} summarizes each work, its approach, and its limitation for the problem addressed here.")
    R.table("lit", "Related systems and papers", ["Work", "Approach", "Limitation or gap"], [
        ["Golem [@golem]", "Peer-to-peer market of requestors and providers, paid in GLM tokens", "Result checking is left to each application; no general availability guarantee for rentals"],
        ["iExec [@iexec]", "Proof-of-Contribution: replicated execution by worker pools, optional TEEs", "Replication multiplies cost; designed for tasks with comparable outputs"],
        ["Akash Network [@akash]", "Reverse auction for container leases on a Cosmos chain, escrow per block", "Payment continues on the provider's word; no on-chain evidence that each block was served"],
        ["Render Network [@render]", "GPU rendering market with node tiers and a job manager", "Job routing and quality checks depend on a central manager"],
        ["Livepeer [@livepeer]", "Staked transcoders with probabilistic micropayments and segment checks", "Specific to video transcoding; checks rely on deterministic output"],
        ["Truebit [@truebit]", "Interactive verification game that resolves disputes over computation", "Needs deterministic, re-executable computation; complex dispute rounds"],
        ["Filecoin [@filecoin]", "Proofs of replication and spacetime for storage, with slashing", "Proofs are specific to storage, not to GPU availability"],
        ["Bittensor [@bittensor]", "Validators score miners on model quality and split token rewards", "Scores are subjective and set by validators; no escrowed rental"],
        ["Arbitrum [@arbitrum]", "Optimistic execution with interactive fraud proofs", "Targets deterministic contract execution, not external hardware"],
        ["Vast.ai [@vast]", "Centralized GPU rental with platform reliability scores", "The platform holds funds and decides disputes"],
    ], [3.2, 6, 6], size=10)
    R.h2("Existing approaches")
    R.p("Three approaches appear in the literature. Re-execution and replication (iExec, Truebit, Livepeer) check that a result is correct by computing it again. Proofs specific to a resource (Filecoin) bind rewards to a cryptographic proof of service. Operator trust (Render, Vast.ai, and in practice many token-based markets) keeps a company in charge of routing, quality, and disputes.")
    R.h2("Limitations and gaps")
    R.bullets([
        "A rented GPU session is interactive and non-deterministic, so it cannot be re-executed and compared.",
        "Lease payments in existing decentralized markets usually continue without per-interval evidence that the provider was online.",
        "Hardware claims are rarely re-checked after registration, so a provider can advertise a GPU it does not serve.",
        "Where slashing exists, it is often tied to a single resource type, and failure costs are not separated from the stake that backs active leases.",
    ])
    R.p("ShaderMesh addresses these gaps with per-interval heartbeats that gate each payout, recurring benchmark challenges with on-chain slashing of free stake, and a separate verification game kept only for tasks that can be recomputed.")

    section(R, "EXISTING SYSTEM")
    R.h2("Current method")
    R.p("Today a person who needs a GPU rents it from a cloud provider or from a centralized marketplace such as Vast.ai. The renter creates an account, pays the platform by card or prepaid balance, and receives access to a machine. The platform pays the host later, after taking a fee. Hardware and reliability scores are kept by the platform, and disputes go to its support team.")
    R.h2("Problems and limitations")
    R.p("{t:existing} lists the main limitations of this model.")
    R.table("existing", "Limitations of the existing system", ["Limitation", "Effect"], [
        ["Central custody of funds", "The platform can freeze, delay, or lose payments; both sides must trust it"],
        ["Opaque reputation", "Hosts and renters cannot verify how scores are computed or what history they reflect"],
        ["No proof of service", "Billing is based on the platform's own records, which the renter cannot check"],
        ["Weak penalties", "A host that goes offline usually loses a rating, not money, so the renter bears the loss"],
        ["Single point of failure", "An outage or policy change at the platform stops the whole market"],
        ["Fees and access barriers", "Platform fees and regional payment restrictions raise the cost of compute"],
    ], [5.5, 9])

    section(R, "PROPOSED SYSTEM")
    R.h2("Proposed solution")
    R.p("ShaderMesh moves the trusted parts of the marketplace into smart contracts and leaves only untrusted coordination off chain. The system offers two market modes over one shared staking, reputation, and slashing layer:")
    R.bullets([
        "**Compute rental (primary).** A requester leases a GPU for a number of fixed intervals and funds escrow for all of them. The provider accepts and bonds part of its stake. For each interval the provider's agent sends a signed heartbeat, and only intervals with a heartbeat are paid. A missed interval ends the lease, refunds the remaining escrow, and slashes the bond.",
        "**Managed inference.** A provider publishes a model offer. A requester opens prepaid credit and sends OpenAI-compatible requests, each signed by the requester's wallet. Usage is settled on chain against the credit.",
        "**Verifiable tasks (retained).** Deterministic jobs, such as rendering frames, are checked by an optimistic verification game with spot checks and slashing, because their output can be recomputed.",
    ])
    R.p("To keep hardware claims honest after registration, an attester issues seeded **benchmark challenges**. The provider must return the result of a 500,000-step SHA-256 hash chain and a measured throughput score within the answer window. A wrong hash, a low score, or a missed deadline slashes 10 SMESH of the provider's free stake to a treasury, revokes its attestation, and lowers its reputation. {f:token} shows how tokens move between the parties.")
    R.figure("token", F / "token-flow.png", "Token flows: deposits, per-interval payouts, refunds, and slashes", max_h=8.0)
    R.h2("Advantages")
    R.bullets([
        "Requester funds are held by a contract, not a company, and are released one interval at a time.",
        "A provider is paid only for intervals it proves with a signed heartbeat, and unused escrow is refunded automatically.",
        "Providers have money at risk: lease bonds for service failures and free stake for failed benchmark challenges.",
        "Reputation, payouts, challenges, and slashes are public events that anyone can audit in the explorer.",
        "Off-chain services hold no requester funds, so a compromised orchestrator cannot steal escrow.",
        "The same layer supports GPU containers, LLM inference, and verifiable batch jobs.",
    ])

    section(R, "BLOCKCHAIN CONCEPTS USED")
    R.p("{t:chain} summarizes the blockchain concepts used in the project.")
    R.table("chain", "Blockchain concepts used", ["Concept", "Choice in ShaderMesh"], [
        ["Blockchain type", "**Public, permissionless.** Ethereum, deployed on the Sepolia test network (chain ID 11155111) and on a local Hardhat chain for development. Anyone can become a provider or requester."],
        ["Consensus mechanism", "**Proof of Stake** [@pos]. Sepolia, like Ethereum mainnet, finalizes blocks through validators that stake ETH. ShaderMesh also uses staking at the application level: providers stake SMESH as collateral."],
        ["Smart contracts", "Six Solidity 0.8.24 contracts built on OpenZeppelin [@oz]: SmeshToken, ProviderRegistry, RentalMarket, InferenceMarket, JobMarket, and Verification. They hold all value and enforce every state transition."],
        ["Token", "**SMESH**, an ERC-20 utility token [@eip20] used for escrow, stake, bonds, credit, and slashing. A faucet issues test tokens on Sepolia. Gas is paid in Sepolia ETH."],
        ["Cryptographic signatures", "Wallet signatures authorize every transaction. EIP-191 signed messages [@eip191] authorize inference requests off chain, and heartbeats are transactions signed by the provider key."],
        ["Hashing", "SHA-256 [@sha2] hash chains for benchmark challenges, keccak256 for challenge seeds and result commitments, and IPFS content identifiers [@ipfs] for workload manifests."],
        ["Events and indexing", "Contracts emit events for every state change. The orchestrator indexes them into PostgreSQL to serve the web application."],
    ], [3.8, 10.7])
    R.p("The contracts are source-verified on Etherscan at the addresses in {t:addr}.")
    R.table("addr", "Sepolia contract addresses", ["Contract", "Address"], [[n, f"`{a}`"] for n, a in ADDR], [4, 10.5], size=10)

    section(R, "SYSTEM ARCHITECTURE AND BLOCK DIAGRAM")
    R.p("{f:arch} shows the containers and the boundary between components that hold value and components that only coordinate. The browser wallet signs every value-moving action. The provider agent runs workloads and sends heartbeats. The attester judges benchmark challenges. The orchestrator indexes events and routes inference, but holds no funds.")
    R.figure("arch", F / "architecture.png", "System architecture: off-chain services and the contracts that hold all value", max_h=11.0)
    R.p("{f:trust} states what each component can and cannot do. The rule behind the design is that a component outside the contracts may fail or lie without being able to move funds it does not own.")
    R.figure("trust", F / "trust-boundary.png", "Trust boundary: what each component can and cannot do", max_h=9.0)

    section(R, "MODULES")
    R.p("The system is split into ten modules, shown in {f:modules}: five on chain and five off chain.")
    R.figure("modules", F / "modules.png", "Module map", max_h=8.0)
    R.h2("On-chain modules")
    R.bullets([
        "**SmeshToken.** ERC-20 token with a rate-limited faucet for test networks.",
        "**ProviderRegistry.** Registers providers with their hardware capabilities, holds stake, records attestations, keeps an integer reputation updated by an exponential moving average, and runs the benchmark challenge state machine (issue, respond, resolve, expire). It slashes only free stake, never stake that backs an active lease.",
        "**RentalMarket.** Holds GPU listings. A lease is requested with escrow and a workload manifest, accepted by the provider with a bond, kept alive by one heartbeat per interval, settled interval by interval, and ended by the requester, by expiry, or by a missed heartbeat.",
        "**InferenceMarket.** Holds model offers and prepaid credit. Usage is settled as cumulative token counts, and the requester can close the credit to recover the unused balance.",
        "**JobMarket and Verification.** Deterministic tasks with bonded claims, result commitments, a challenge window, and a vote that slashes the losing side.",
    ])
    R.h2("Off-chain modules")
    R.bullets([
        "**Provider agent (Python).** Discovers the GPU, accepts allowed lease manifests, runs workloads in constrained Docker containers with GPU passthrough, serves inference through Ollama, sends heartbeats, and answers benchmark challenges.",
        "**Attester (Python).** Issues a challenge to each attested provider at a fixed interval, checks the hash and score, and records the result on chain. Anyone can expire a challenge that was not answered in time.",
        "**Orchestrator (TypeScript).** Indexes contract events into PostgreSQL, serves REST and server-sent events to the web application, and routes signed inference requests to the provider that holds the offer.",
        "**Web application (Next.js).** A landing page and an app with three workspaces: Rent (Marketplace, My leases, Inference), Provide (Node, Capacity, Stake), and an Explorer for leases, challenges, and events.",
        "**Simulator (Python).** A Monte Carlo model of the economics, used to choose spot-check rates and to check that rental escrow is always conserved.",
    ])
    R.p("{f:ui-home} and {f:ui-rent} show the web application running against the live Sepolia deployment.")
    R.figure("ui-home", S / "dashboard.png", "Web application: Home page with the Renting and Providing panels", max_h=8.5)
    R.figure("ui-rent", S / "rent.png", "Web application: Marketplace with the lease composer", max_h=8.5)

    section(R, "TECHNOLOGY STACK")
    R.p("{f:stack} groups the technologies by layer, and {t:stack} lists them under the headings of the template.")
    R.figure("stack", F / "tech-stack.png", "Technology stack by layer", max_h=7.5)
    R.table("stack", "Technology stack", ["Category", "Technology"], [
        ["Blockchain platform", "Ethereum (Sepolia test network), local Hardhat network for development"],
        ["Programming languages", "Solidity 0.8.24 (contracts), Python 3.11 (agent, attester, simulator), TypeScript (orchestrator, web)"],
        ["Database and storage", "PostgreSQL 16 (event projection), IPFS (workload manifests and job data)"],
        ["Frameworks and libraries", "OpenZeppelin, Hardhat [@hardhat], Next.js 16 [@next], React 19, Tailwind CSS, viem [@viem], ethers, web3.py [@web3py]"],
        ["Compute runtime", "Docker Engine with NVIDIA GPU passthrough [@docker], Ollama [@ollama] for LLM inference"],
        ["Wallet", "MetaMask through the EIP-1193 provider interface [@eip1193]"],
        ["Testing and delivery", "Mocha and Chai, pytest, Playwright [@playwright], Docker Compose, GitHub Actions"],
    ], [4.3, 10.2])

    section(R, "METHODOLOGY AND WORKFLOW")
    R.h2("Development methodology")
    R.p("The project was built in seven phases ({f:method}). Each phase delivered working code with automated tests before the next one started, and findings from testing were fed back into the requirements. Contracts came first because every other component depends on their interfaces and events.")
    R.figure("method", F / "methodology.png", "Development phases", max_h=5.0)
    R.h2("Rental workflow")
    R.p("{f:rental} shows the main user workflow, a GPU rental from listing to settlement. The steps are:")
    R.numbered([
        "The provider registers, stakes SMESH, and is attested by the attester. It then publishes a listing with a price per interval.",
        "The requester chooses a GPU in the Marketplace, sets the number of intervals and a workload manifest, approves the escrow, and requests the lease.",
        "The provider agent checks the manifest against its allowlist and accepts the lease, which bonds part of its stake and starts the clock.",
        "The agent runs the workload in a constrained container and sends one heartbeat transaction per interval.",
        "Each interval with a heartbeat can be settled, which pays the provider. If an interval has no heartbeat, the lease ends: the remaining escrow is refunded and the bond is slashed.",
        "When the lease ends normally, the bond returns to the provider's free stake and any unused escrow returns to the requester.",
    ])
    R.figure("rental", F / "rental-sequence.png", "Rental workflow: listing, lease, heartbeats, settlement, and the two ways a lease ends", max_h=11.0)
    R.h2("Benchmark challenge workflow")
    R.p("{f:challenge} shows the benchmark challenge. The attester issues a challenge with a seed derived from the block's randomness, the provider computes the hash chain and measures throughput, and the attester judges the answer. A challenge that is not answered or not judged within the window can be expired by anyone.")
    R.figure("challenge", F / "challenge-state-machine.png", "Benchmark challenge states and transitions", max_h=7.0)

    section(R, "DATASET AND INPUT DETAILS")
    R.p("The project does not train a model and does not use a fixed dataset. Its inputs are produced by users and by the network, as listed in {t:inputs}.")
    R.table("inputs", "System inputs", ["Input", "Source", "Format and use"], [
        ["Hardware descriptor", "Provider agent", "GPU model, VRAM, region, and models served; registered on chain as capabilities"],
        ["Workload manifest", "Requester", "IPFS content identifier of a container manifest; bound to the lease and checked against the provider's allowlist"],
        ["Inference prompt", "Requester", "OpenAI-compatible chat request, signed with EIP-191 and carrying a unique request ID"],
        ["Challenge seed", "Contract", "keccak256 of block randomness, provider address, and challenge ID; input to the SHA-256 hash chain"],
        ["Lease parameters", "Requester and provider", "Price per interval, interval length, number of intervals, and bond"],
        ["Simulator parameters", "Project team", "Reward, bond, spot-check rate, challenger pool, outage rates; used to evaluate the economics"],
    ], [3.6, 3.2, 7.7], size=10.5)

    section(R, "EXPECTED OUTCOMES")
    R.p("{t:outcomes} lists the expected outcomes with the status reached at the time of this report.")
    R.table("outcomes", "Expected outcomes and status", ["Outcome", "Status"], [
        ["Contracts for escrow, stake, rental, inference, verifiable tasks, and benchmark challenges", "Done: six contracts, 42 passing tests"],
        ["Provider agent that runs GPU containers and LLM inference", "Done: 58 passing tests; RTX 4070 Laptop GPU serving qwen3:0.6b"],
        ["Indexer and API that can be rebuilt from chain events", "Done: 54 passing tests with PostgreSQL"],
        ["Web application for requesters and providers", "Done: role workspaces; 30 browser checks including wallet transactions"],
        ["Deployment on a public test network", "Done: all six contracts deployed and source-verified on Sepolia"],
        ["Benchmark challenges on a live provider", "Done: challenges answered and passed on Sepolia, about 289 tokens per second measured"],
        ["Evaluation of cost and economics", "Done: gas report for 27 transactions; simulator with 15 passing tests"],
    ], [8.5, 6])
    R.p("Measured gas for the main rental and challenge transactions is shown in {t:gas}. At typical test-network gas prices these costs are small compared to the value of an hour of GPU time.")
    R.table("gas", "Gas used by core transactions", ["Transaction", "Gas used"], [
        ["Request lease", "251,540"], ["Accept lease", "154,207"], ["Heartbeat (per interval)", "56,487"],
        ["Settle interval", "59,987"], ["Issue benchmark challenge", "129,754"], ["Respond to challenge", "76,217"],
        ["Resolve challenge (pass)", "54,309"], ["Expire challenge (slash)", "80,193"], ["Open inference credit", "227,898"],
    ], [10, 4.5], align_right=(1,))
    R.p("The simulation reproduces the expected break-even spot-check rate of 0.5 when reward and bond are equal, and confirms that rental escrow is always conserved: every token deposited is either paid, refunded, or slashed ({f:econ}).")
    R.figure("econ", F / "chart-rental-economics.png", "Simulated rental outcomes by provider type: paid, refunded, and slashed", max_h=7.5)
    R.p("Known limits remain. The attester is a trusted role, LLM output quality is not verified on chain, and a heartbeat proves that the provider's key was live, not the quality of service. These are documented as future work.")


def main():
    def build(labels):
        R = Report(REFS, labels)
        h1 = R.d.styles["Heading 1"].paragraph_format
        h1.page_break_before = False
        h1.alignment = AL.LEFT
        h1.space_before = R.d.styles["Heading 2"].paragraph_format.space_before
        R.d.core_properties.title = "ShaderMesh: Project Report"
        cover(R)
        body(R)
        R.plain_h1 = lambda t: section(R, t)
        R.references()
        return R

    first = build(None)
    final = build(first.new_labels)
    final.d.save(str(OUT_DOCX))
    print("saved", OUT_DOCX, "figures/tables:", len(final.new_labels))
    if "--no-pdf" not in sys.argv:
        print("saved", OUT_PDF, "total pages", export_pdf(OUT_DOCX, OUT_PDF))


if __name__ == "__main__":
    main()
